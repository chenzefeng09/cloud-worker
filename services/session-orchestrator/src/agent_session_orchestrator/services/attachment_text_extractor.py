"""Best-effort text extraction from user-uploaded attachment bytes.

Used by the inline (chat-mode) run executor so a model can *read* files the
user attached to the session — PDF / Word / Excel / PowerPoint / plain-text
formats are decoded to plain text and injected into the system prompt.

Design contract (the whole chain is fail-open — a broken/oversized/unknown
file must NEVER crash a run):

* :func:`extract_text` dispatches on mime + filename extension and returns a
  plain-text string.  It *never raises*: any parse error is caught and turned
  into a short Chinese diagnostic string (``（文件解析失败：…）``).
* Unsupported binary formats (legacy .doc/.ppt/.xls, images, archives) return a
  fixed Chinese "unsupported" notice rather than garbage.
* Every extractor is size- and length-bounded (see the module constants) so a
  pathological file can't blow up the prompt or the event loop.

The heavy parsing libraries (``pypdf`` / ``python-docx`` / ``openpyxl`` /
``python-pptx``) are imported lazily inside each branch so importing this module
stays cheap and a missing optional dependency degrades to the fail-open notice
instead of an ImportError at module load.
"""

from __future__ import annotations

import io
import logging

logger = logging.getLogger(__name__)

# ── Limits ──────────────────────────────────────────────────────────────────
# Per-file extracted-text ceiling (characters).  Beyond this the text is cut and
# the truncation marker appended.
MAX_CHARS_PER_FILE = 30_000
# Whole-session extraction budget (characters) shared across all attachments;
# enforced by the caller (:func:`build_attachments_block`).
MAX_CHARS_PER_SESSION = 60_000
# Raw-byte ceiling per file: above this we skip parsing entirely (a 50 MB PDF
# would take too long / use too much memory to be worth reading inline).
MAX_BYTES_PER_FILE = 20 * 1024 * 1024
# Per-sheet row cap for spreadsheets — a huge sheet is truncated, not dumped.
MAX_XLSX_ROWS_PER_SHEET = 200

_TRUNCATED_MARK = "\n……（内容过长，已截断）"
_UNSUPPORTED_MSG = "（暂不支持解析该格式）"

# Appended to the system prompt whenever an attachment / workspace-file block was
# injected.  Instructs the model to reference those files with a markdown link
# whose target is the file id given in the block, which the frontend intercepts to
# open a preview panel.  Only files actually listed above carry an id — the model
# must never invent one.
FILE_LINK_CONVENTION_PROMPT = (
    "当你在回复中提到上述文件时，请用 markdown 链接 [文件名](#file:文件id) 的形式引用该文件，"
    "其中「文件id」必须使用对应文件块中 id 属性标注的真实 id（形如 att_xxx 或 wsf_xxx）。"
    "只对上面真实列出的文件使用这种格式，不要编造 id，也不要给未列出的文件加链接。"
    "用户点击该链接即可直接预览对应文件。"
)


def truncate_text(text: str, limit: int = MAX_CHARS_PER_FILE) -> str:
    """Trim ``text`` to ``limit`` chars, appending the truncation marker."""
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + _TRUNCATED_MARK


# Backward-compatible private alias for callers/tests written before the rich
# document converter needed the same public prompt-size guard.
_truncate = truncate_text


def _ext(filename: str) -> str:
    """Lower-cased file extension including the dot (``.pdf``), or ``""``."""
    name = (filename or "").strip().lower()
    dot = name.rfind(".")
    return name[dot:] if dot >= 0 else ""


# ── Format detectors ─────────────────────────────────────────────────────────
_TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".yaml",
    ".yml", ".xml", ".html", ".htm", ".py", ".js", ".ts", ".java", ".go",
    ".rs", ".c", ".h", ".cpp", ".sh", ".sql", ".ini", ".toml", ".rst",
}
# Legacy binary Office formats + images + archives we deliberately do not parse.
_UNSUPPORTED_EXTS = {
    ".doc", ".ppt", ".xls", ".rtf",
    ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tiff", ".svg", ".ico",
    ".zip", ".rar", ".7z", ".tar", ".gz", ".bz2",
    ".mp3", ".mp4", ".wav", ".avi", ".mov", ".mkv",
    ".bin", ".exe", ".dll", ".so", ".dylib",
}


def _looks_text(mime: str, ext: str) -> bool:
    m = (mime or "").lower()
    if m.startswith("text/"):
        return True
    if m in {
        "application/json", "application/xml", "application/x-yaml",
        "application/yaml", "application/csv", "application/markdown",
    }:
        return True
    return ext in _TEXT_EXTS


def _is_pdf(mime: str, ext: str) -> bool:
    return ext == ".pdf" or (mime or "").lower() == "application/pdf"


def _is_docx(mime: str, ext: str) -> bool:
    m = (mime or "").lower()
    return ext == ".docx" or m == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )


def _is_xlsx(mime: str, ext: str) -> bool:
    m = (mime or "").lower()
    return ext == ".xlsx" or m == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def _is_pptx(mime: str, ext: str) -> bool:
    m = (mime or "").lower()
    return ext == ".pptx" or m == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )


# ── Per-format extractors (each returns plain text, never raises) ────────────
def _extract_text_plain(data: bytes) -> str:
    """Decode text bytes: utf-8 → gbk → utf-8/replace (never fails)."""
    for enc in ("utf-8", "gbk"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader  # noqa: PLC0415

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 — a bad page shouldn't kill the file
            page_text = ""
        if page_text.strip():
            parts.append(page_text.strip())
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    import docx  # noqa: PLC0415  (python-docx)

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    for para in document.paragraphs:
        if para.text and para.text.strip():
            parts.append(para.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet_texts: list[str] = []
        for ws in wb.worksheets:
            rows: list[str] = []
            truncated = False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_XLSX_ROWS_PER_SHEET:
                    truncated = True
                    break
                cells = ["" if v is None else str(v) for v in row]
                # Skip fully blank rows to keep the dump compact.
                if any(c.strip() for c in cells):
                    rows.append("\t".join(cells))
            body = "\n".join(rows)
            if truncated:
                body += f"\n……（工作表行数超过 {MAX_XLSX_ROWS_PER_SHEET} 行，已截断）"
            sheet_texts.append(f"# 工作表：{ws.title}\n{body}")
        return "\n\n".join(sheet_texts)
    finally:
        wb.close()


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation  # noqa: PLC0415

    prs = Presentation(io.BytesIO(data))
    slide_texts: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if text:
                        lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        lines.append("\t".join(cells))
        body = "\n".join(lines)
        slide_texts.append(f"第{idx}页：\n{body}" if body else f"第{idx}页：（无文本）")
    return "\n\n".join(slide_texts)


def extract_text(filename: str, mime: str, data: bytes) -> str:
    """Extract plain text from an attachment's raw bytes.

    Dispatches on mime + extension.  Fail-open on every path: an unsupported
    format returns a fixed Chinese notice, a parse error returns
    ``（文件解析失败：…）``, an over-size file is skipped with a notice.  The
    returned text is length-bounded to :data:`MAX_CHARS_PER_FILE`.
    """
    if not data:
        return "（文件为空）"
    if len(data) > MAX_BYTES_PER_FILE:
        mb = MAX_BYTES_PER_FILE // (1024 * 1024)
        return f"（文件过大，超过 {mb}MB，已跳过解析）"

    ext = _ext(filename)

    # Explicit unsupported binary formats short-circuit before any parse attempt.
    if ext in _UNSUPPORTED_EXTS:
        return _UNSUPPORTED_MSG

    try:
        if _is_pdf(mime, ext):
            text = _extract_pdf(data)
        elif _is_docx(mime, ext):
            text = _extract_docx(data)
        elif _is_xlsx(mime, ext):
            text = _extract_xlsx(data)
        elif _is_pptx(mime, ext):
            text = _extract_pptx(data)
        elif _looks_text(mime, ext):
            text = _extract_text_plain(data)
        else:
            return _UNSUPPORTED_MSG
    except Exception as exc:  # noqa: BLE001 — fail-open, never crash a run
        logger.warning(
            "attachment text extraction failed",
            extra={"att_filename": filename, "att_mime": mime, "err": str(exc)},
        )
        return f"（文件解析失败：{exc}）"

    text = text.strip()
    if not text:
        return "（未从文件中解析出文本内容）"
    return _truncate(text)


def build_attachments_block(
    files: list[tuple[str, str, str]],
    *,
    session_budget: int = MAX_CHARS_PER_SESSION,
) -> str | None:
    """Assemble the Chinese system-prompt block from ``(id, filename, text)`` triples.

    ``files`` is a list of already-extracted ``(attachment_id, filename,
    extracted_text)`` triples.  The ``attachment_id`` (``att_…``) is surfaced as
    an ``id`` attribute on each ``<attachment>`` tag so the model can reference a
    file back to the caller via the ``[文件名](#file:<id>)`` link convention
    (see :data:`FILE_LINK_CONVENTION_PROMPT`).  Returns the full injection block,
    or ``None`` when there is nothing to inject (empty input).  Enforces the
    whole-session character budget: once the accumulated text would exceed
    ``session_budget`` the remaining files are dropped with a notice.
    """
    if not files:
        return None

    header = (
        "用户在本会话上传了以下文件，内容已为你解析如下。"
        "回答与这些文件相关的问题时请直接引用其内容，不要声称无法读取文件。"
    )
    parts: list[str] = [header]
    used = 0
    dropped = 0
    for file_id, filename, text in files:
        body = text or ""
        remaining = session_budget - used
        if remaining <= 0:
            dropped += 1
            continue
        if len(body) > remaining:
            body = body[:remaining].rstrip() + _TRUNCATED_MARK
        used += len(body)
        # Escape nothing — the model reads this as context, and the filename is
        # rendered as an attribute for readability only.
        safe_name = (filename or "未命名文件").replace('"', "'")
        # The id attribute is emitted only when a real id is present so the model
        # is never handed an empty/fabricated anchor (see the link convention).
        id_attr = f' id="{file_id}"' if file_id else ""
        parts.append(
            f'<attachment{id_attr} filename="{safe_name}">\n{body}\n</attachment>'
        )
    if dropped:
        parts.append(f"（另有 {dropped} 个文件因内容预算已满未展示）")

    return "\n\n".join(parts)


# ── Workspace base files ─────────────────────────────────────────────────────
# The workspace-files block gets its own, separate budget so a heavyweight
# workspace can never starve the session-attachment block (and vice versa).
MAX_CHARS_PER_WORKSPACE = 30_000
MAX_WORKSPACE_FILES = 10


def build_workspace_files_block(
    files: list[tuple[str, str, str]],
    *,
    session_budget: int = MAX_CHARS_PER_WORKSPACE,
    max_files: int = MAX_WORKSPACE_FILES,
) -> str | None:
    """Assemble the workspace base-files system-prompt block.

    Same contract as :func:`build_attachments_block` — ``files`` is a list of
    already-extracted ``(workspace_file_id, filename, text)`` triples, the return
    is the full injection block or ``None`` when there is nothing to inject — but
    with a ``<workspace_file>`` tag, an independent character budget and a hard
    file cap.  The ``workspace_file_id`` (``wsf_…``) is surfaced as an ``id``
    attribute so the model can reference the file via the
    ``[文件名](#file:<id>)`` link convention (see
    :data:`FILE_LINK_CONVENTION_PROMPT`).  Files beyond either limit are not
    injected; their names are listed at the end so the model (and the user
    reading a prompt dump) knows what was omitted.
    """
    if not files:
        return None

    header = (
        "当前会话所属工作区配置了以下基础文件，内容已为你解析如下。"
        "回答与这些文件相关的问题时请直接引用其内容，不要声称无法读取文件。"
    )
    parts: list[str] = [header]
    used = 0
    omitted: list[str] = []
    for index, (file_id, filename, text) in enumerate(files):
        name = filename or "未命名文件"
        if index >= max_files:
            omitted.append(name)
            continue
        body = text or ""
        remaining = session_budget - used
        if remaining <= 0:
            omitted.append(name)
            continue
        if len(body) > remaining:
            body = body[:remaining].rstrip() + _TRUNCATED_MARK
        used += len(body)
        safe_name = name.replace('"', "'")
        # id attribute only when present (never hand the model an empty anchor).
        id_attr = f' id="{file_id}"' if file_id else ""
        parts.append(
            f'<workspace_file{id_attr} filename="{safe_name}">\n{body}\n</workspace_file>'
        )
    if omitted:
        parts.append(
            "（另有以下工作区文件因数量或内容预算上限未注入："
            + "、".join(omitted)
            + "）"
        )

    return "\n\n".join(parts)
